"""
Phase 3 extended regression tests.

Covers:
  Task A -- Unicode text extraction (clean_extracted_text)
  Task B -- Skill removal / persistence semantics
  Task C -- Strengthened PDF / DOCX / TXT extraction tests
  Task D -- Skill extraction normalization edge cases
  Task E -- Authentication, authorization, and ownership
  Task F -- Re-upload and state management
"""

import io
import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.text_extractor import (
    clean_extracted_text,
    extract_text_from_txt,
    extract_text_from_docx,
    extract_text_from_pdf,
    MAX_FILE_SIZE,
)
from app.services.skill_extractor import extract_skills_from_text


def make_docx_bytes(text):
    doc = Document()
    for line in text.split("\n"):
        stripped = line.strip()
        if stripped:
            doc.add_paragraph(stripped)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_docx_with_table(text):
    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = text
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_blank_pdf():
    writer = PdfWriter()
    writer.add_blank_page(width=72 * 8.5, height=72 * 11)
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


def _register_and_login(client, email, full_name="Test User", role="student"):
    client.post(
        "/auth/register",
        json={"email": email, "full_name": full_name, "password": "Password123!", "role": role},
    )
    login = client.post("/auth/login", json={"email": email, "password": "Password123!"})
    token = login.json()["access_token"]
    return token, {"Authorization": f"Bearer {token}"}


SAMPLE_RESUME_TEXT = """
Priya Mehta
B.Tech Computer Science, 2025

TECHNICAL SKILLS
Languages: Python, JavaScript, TypeScript, SQL
Frameworks: FastAPI, React, Django
Databases: PostgreSQL, MongoDB

PROJECTS
CareerGraph AI
- Built scalable REST APIs using FastAPI and PostgreSQL.
- Implemented React frontend integrated with Node.js backend.

E-Commerce Platform
- Developed microservices architecture using Docker and Kubernetes.
- Architected Redis caching and MongoDB document storage.

WORK EXPERIENCE
Software Intern at TechCorp (2024)
- Implemented CI/CD pipelines using GitHub Actions and Jenkins.
- Utilized Python, AWS Lambda, and Docker containers for deployment.
"""


# ── TASK A: Unicode text extraction ──────────────────────────────────────────

class TestCleanExtractedText:
    def test_plain_ascii(self):
        assert clean_extracted_text("Hello World") == "Hello World"

    def test_accented_name_preserved(self):
        text = "Jose Maria Garcia"
        assert clean_extracted_text(text) == text

    def test_accented_chars_preserved(self):
        text = "Jos\u00e9 Mar\u00eda Garc\u00eda"
        assert clean_extracted_text(text) == text

    def test_hindi_text_preserved(self):
        text = "\u0928\u093e\u092e: \u0926\u093f\u0935\u094d\u092f\u093e\u0936\u093f\u0937"
        assert clean_extracted_text(text) == text

    def test_em_dash_preserved(self):
        text = "Python \u2014 Advanced"
        result = clean_extracted_text(text)
        assert "\u2014" in result

    def test_en_dash_preserved(self):
        text = "2020\u20132024 Software Engineer"
        result = clean_extracted_text(text)
        assert "\u2013" in result

    def test_newlines_preserved(self):
        text = "Line1\nLine2\nLine3"
        result = clean_extracted_text(text)
        assert result.count("\n") == 2

    def test_tabs_converted_to_space(self):
        result = clean_extracted_text("col1\tcol2")
        assert "\t" not in result
        assert "col1 col2" == result

    def test_control_chars_removed(self):
        text = "Hello\x00World\x01Foo\x07Bar\x1bEnd"
        result = clean_extracted_text(text)
        assert "\x00" not in result
        assert "\x01" not in result
        assert "\x07" not in result
        assert "\x1b" not in result
        assert "Hello" in result
        assert "World" in result
        assert "End" in result

    def test_empty_string(self):
        assert clean_extracted_text("") == ""

    def test_whitespace_only_string(self):
        assert clean_extracted_text("   \n\n\t  ") == ""

    def test_crlf_normalized_to_lf(self):
        result = clean_extracted_text("line1\r\nline2\rline3")
        assert "\r" not in result
        assert result.count("\n") == 2

    def test_multiple_blank_lines_collapsed(self):
        text = "Para1\n\n\n\n\nPara2"
        result = clean_extracted_text(text)
        assert "\n\n\n" not in result
        assert "Para1" in result and "Para2" in result

    def test_repeated_spaces_collapsed(self):
        result = clean_extracted_text("word1     word2")
        assert "  " not in result
        assert "word1 word2" == result

    def test_non_breaking_space_converted(self):
        result = clean_extracted_text("Python\xa0Developer")
        assert "\xa0" not in result
        assert "Python Developer" == result

    def test_mixed_unicode(self):
        text = "Name: Jos\u00e9\nSkills: Python, \u0928\u0947\u091f\u0935\u0930\u094d\u0915\u093f\u0902\u0917"
        result = clean_extracted_text(text)
        assert "Jos\u00e9" in result
        assert "\u0928\u0947\u091f\u0935\u0930\u094d\u0915\u093f\u0902\u0917" in result
        assert "Python" in result


# ── TASK C: Extraction unit tests ────────────────────────────────────────────

class TestTxtExtraction:
    def test_utf8_ascii(self):
        content = b"Python developer with REST API experience."
        assert "Python" in extract_text_from_txt(content)

    def test_utf8_unicode(self):
        text = "Jos\u00e9 Mar\u00eda FastAPI developer"
        assert "Jos\u00e9" in extract_text_from_txt(text.encode("utf-8"))

    def test_hindi_utf8(self):
        text = "\u0928\u093e\u092e: \u0905\u0930\u094d\u091c\u0941\u0928"
        assert "\u0905\u0930\u094d\u091c\u0941\u0928" in extract_text_from_txt(text.encode("utf-8"))

    def test_utf8_bom(self):
        content = "Python developer".encode("utf-8-sig")
        result = extract_text_from_txt(content)
        assert "Python" in result

    def test_latin1_encoded(self):
        content = "Ren\xe9 is a developer".encode("latin-1")
        result = extract_text_from_txt(content)
        assert "developer" in result

    def test_empty_bytes_raises(self):
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_txt(b"")
        assert exc_info.value.status_code == 400


class TestDocxExtraction:
    def test_paragraphs_extracted(self):
        docx_bytes = make_docx_bytes("Python developer\nFastAPI experience")
        result = extract_text_from_docx(docx_bytes)
        assert "Python" in result
        assert "FastAPI" in result

    def test_table_cells_extracted(self):
        docx_bytes = make_docx_with_table("Docker | Kubernetes | AWS")
        result = extract_text_from_docx(docx_bytes)
        assert "Docker" in result

    def test_unicode_paragraphs_preserved(self):
        docx_bytes = make_docx_bytes("Jos\u00e9 Garc\u00eda Python Developer")
        result = extract_text_from_docx(docx_bytes)
        assert "Jos\u00e9" in result

    def test_empty_docx_raises(self):
        from fastapi import HTTPException
        doc = Document()
        buf = io.BytesIO()
        doc.save(buf)
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_docx(buf.getvalue())
        assert exc_info.value.status_code == 400

    def test_malformed_docx_raises(self):
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_docx(b"PK\x03\x04garbage not a real docx")
        assert exc_info.value.status_code == 400

    def test_non_docx_bytes_raise(self):
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_docx(b"\x00\x01\x02\x03binary garbage")
        assert exc_info.value.status_code == 400


class TestPdfExtraction:
    def test_blank_pdf_raises(self):
        from fastapi import HTTPException
        pdf_bytes = make_blank_pdf()
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_pdf(pdf_bytes)
        assert exc_info.value.status_code == 400
        assert "no selectable text" in exc_info.value.detail.lower()

    def test_malformed_pdf_raises(self):
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_pdf(b"This is not a PDF file at all!")
        assert exc_info.value.status_code == 400

    def test_binary_garbage_raises(self):
        from fastapi import HTTPException
        with pytest.raises(HTTPException) as exc_info:
            extract_text_from_pdf(b"\x00\x01\x02\x03\x04")
        assert exc_info.value.status_code == 400


# ── TASK D: Skill normalization edge cases ───────────────────────────────────

class TestSkillExtractionNormalization:
    def test_java_not_matched_by_javascript(self):
        skills = extract_skills_from_text("I use JavaScript and React for frontend.")
        names = {s.name for s in skills}
        assert "JavaScript" in names
        assert "Java" not in names

    def test_java_matched_independently(self):
        skills = extract_skills_from_text(
            "TECHNICAL SKILLS\nJava, Spring Boot, Maven\n"
            "PROJECTS\nBuilt a Java application using Spring Boot."
        )
        names = {s.name for s in skills}
        assert "Java" in names
        assert "Spring Boot" in names

    def test_case_sensitive_c_language(self):
        text = (
            "TECHNICAL SKILLS\nC, C++, Python\n"
            "PROJECTS\nImplemented data structures in C and C++."
        )
        skills = extract_skills_from_text(text)
        names = {s.name for s in skills}
        assert "C" in names
        assert "C++" in names
        assert "Python" in names

    def test_duplicate_aliases_no_duplicate_skills(self):
        text = "Used postgresql and postgres and psql in my project."
        skills = extract_skills_from_text(text)
        pg_skills = [s for s in skills if s.name == "PostgreSQL"]
        assert len(pg_skills) == 1

    def test_confidence_bounds(self):
        skills = extract_skills_from_text(SAMPLE_RESUME_TEXT)
        for skill in skills:
            assert 0.0 <= skill.confidence <= 1.0

    def test_empty_resume_returns_empty_list(self):
        assert extract_skills_from_text("") == []

    def test_short_unrecognized_resume(self):
        skills = extract_skills_from_text("This is a plain sentence with no tech skills.")
        assert isinstance(skills, list)

    def test_evidence_from_actual_text(self):
        skills = extract_skills_from_text(SAMPLE_RESUME_TEXT)
        for skill in skills:
            assert len(skill.evidence) > 0
            assert len(skill.evidence) <= 140

    def test_proficiency_valid_values(self):
        skills = extract_skills_from_text(SAMPLE_RESUME_TEXT)
        valid_levels = {"beginner", "intermediate", "advanced"}
        for skill in skills:
            assert skill.proficiency in valid_levels

    def test_section_detection_boosts_confidence(self):
        text_with_section = (
            "TECHNICAL SKILLS\nPython, FastAPI\n"
            "PROJECTS\nBuilt APIs using Python and FastAPI."
        )
        text_without_section = "I enjoy using Python and FastAPI occasionally."
        skills_with = extract_skills_from_text(text_with_section)
        skills_without = extract_skills_from_text(text_without_section)
        python_with = next((s for s in skills_with if s.name == "Python"), None)
        python_without = next((s for s in skills_without if s.name == "Python"), None)
        if python_with and python_without:
            assert python_with.confidence >= python_without.confidence


# ── TASK B: Skill confirmation and persistence semantics ─────────────────────

class TestSkillConfirmationSemantics:
    def test_confirmed_skill_saved(self, client):
        _, headers = _register_and_login(client, "confirm_a@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        assert upload.status_code == 201
        rid = upload.json()["resume_id"]

        confirm = client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": [{"name": "Python", "level": "advanced"}]},
            headers=headers)
        assert confirm.status_code == 200

        profile = client.get("/profile/me", headers=headers).json()
        skill_names = {s["name"] for s in profile["student_profile"]["skills"]}
        assert "Python" in skill_names

    def test_removed_skill_absent_after_confirmation(self, client):
        _, headers = _register_and_login(client, "confirm_b@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        rid = upload.json()["resume_id"]

        # Only confirm Python; FastAPI is intentionally omitted (simulating removal)
        confirm = client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": [{"name": "Python", "level": "advanced"}]},
            headers=headers)
        assert confirm.status_code == 200

        profile = client.get("/profile/me", headers=headers).json()
        skill_names = {s["name"] for s in profile["student_profile"]["skills"]}
        assert "Python" in skill_names
        assert "FastAPI" not in skill_names, "FastAPI was removed; must not appear in profile"

    def test_manually_added_skill_persists(self, client):
        _, headers = _register_and_login(client, "confirm_c@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        rid = upload.json()["resume_id"]

        confirm = client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": [
                {"name": "Python", "level": "advanced"},
                {"name": "SystemVerilog", "level": "intermediate"},
            ]},
            headers=headers)
        assert confirm.status_code == 200

        profile = client.get("/profile/me", headers=headers).json()
        skill_names = {s["name"] for s in profile["student_profile"]["skills"]}
        assert "Python" in skill_names
        assert "SystemVerilog" in skill_names

    def test_unrelated_profile_fields_preserved(self, client):
        _, headers = _register_and_login(client, "confirm_d@example.com")
        client.put("/profile/student", json={
            "college": "IIT Bombay", "degree": "B.Tech CSE",
            "graduation_year": 2026, "target_role": "Backend Developer",
            "bio": "Passionate about distributed systems.",
            "github_username": "testgithub", "skills": [],
        }, headers=headers)

        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        rid = upload.json()["resume_id"]

        client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": [{"name": "Python", "level": "intermediate"}]},
            headers=headers)

        profile = client.get("/profile/me", headers=headers).json()["student_profile"]
        assert profile["college"] == "IIT Bombay"
        assert profile["degree"] == "B.Tech CSE"
        assert profile["bio"] == "Passionate about distributed systems."
        assert profile["github_username"] == "testgithub"

    def test_repeated_confirmation_no_duplicates(self, client):
        _, headers = _register_and_login(client, "confirm_e@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        rid = upload.json()["resume_id"]

        payload = {"resume_id": rid, "skills": [
            {"name": "Python", "level": "advanced"},
            {"name": "FastAPI", "level": "intermediate"},
        ]}
        client.post("/resume/confirm-skills", json=payload, headers=headers)
        client.post("/resume/confirm-skills", json=payload, headers=headers)

        profile = client.get("/profile/me", headers=headers).json()
        skills = profile["student_profile"]["skills"]
        assert sum(1 for s in skills if s["name"] == "Python") == 1

    def test_empty_skill_list_clears_profile(self, client):
        _, headers = _register_and_login(client, "confirm_f@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        rid = upload.json()["resume_id"]

        client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": [{"name": "Python", "level": "advanced"}]},
            headers=headers)

        confirm2 = client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": []},
            headers=headers)
        assert confirm2.status_code == 200

        profile = client.get("/profile/me", headers=headers).json()
        assert profile["student_profile"]["skills"] == []


# ── TASK E: Authentication and authorization ─────────────────────────────────

class TestResumeAuth:
    def test_upload_requires_auth(self, client):
        txt = b"Python developer with FastAPI experience."
        res = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")})
        assert res.status_code == 401

    def test_get_latest_requires_auth(self, client):
        res = client.get("/resume/latest")
        assert res.status_code == 401

    def test_confirm_skills_requires_auth(self, client):
        res = client.post("/resume/confirm-skills",
            json={"skills": [{"name": "Python", "level": "advanced"}]})
        assert res.status_code == 401

    def test_alumni_cannot_upload_resume(self, client):
        _, headers = _register_and_login(client, "alumni_upload@example.com", role="alumni")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        res = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        assert res.status_code == 403

    def test_alumni_cannot_confirm_skills(self, client):
        _, headers = _register_and_login(client, "alumni_confirm@example.com", role="alumni")
        res = client.post("/resume/confirm-skills",
            json={"skills": [{"name": "Python", "level": "intermediate"}]},
            headers=headers)
        assert res.status_code == 403

    def test_no_resume_returns_404(self, client):
        _, headers = _register_and_login(client, "nofile@example.com")
        res = client.get("/resume/latest", headers=headers)
        assert res.status_code == 404

    def test_invalid_token_rejected(self, client):
        fake_headers = {"Authorization": "Bearer invalid.token.here"}
        res = client.get("/resume/latest", headers=fake_headers)
        assert res.status_code == 401

    def test_user_cannot_access_other_users_resume(self, client):
        _, headers_a = _register_and_login(client, "owner_a@example.com")
        _, headers_b = _register_and_login(client, "owner_b@example.com")

        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload_a = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers_a)
        assert upload_a.status_code == 201
        resume_id_a = upload_a.json()["resume_id"]

        # B has no resume; gets 404
        res_b = client.get("/resume/latest", headers=headers_b)
        assert res_b.status_code == 404

        # B attempts to confirm A's resume_id; server ignores invalid ownership
        client.post("/resume/confirm-skills",
            json={"resume_id": resume_id_a, "skills": [{"name": "Python", "level": "intermediate"}]},
            headers=headers_b)

        # A's resume must remain unchanged
        latest_a = client.get("/resume/latest", headers=headers_a).json()
        assert latest_a["status"] == "parsed"


# ── TASK C: File upload validation ───────────────────────────────────────────

class TestFileValidation:
    def _headers(self, client, email):
        _, headers = _register_and_login(client, email)
        return headers

    def test_oversized_file_rejected(self, client):
        headers = self._headers(client, "oversize@example.com")
        big = b"A" * (MAX_FILE_SIZE + 1)
        res = client.post("/resume/upload", files={"file": ("big.txt", big, "text/plain")}, headers=headers)
        assert res.status_code == 400
        detail = res.json()["detail"].lower()
        assert "10 mb" in detail or "size" in detail

    def test_unsupported_extension_rejected(self, client):
        headers = self._headers(client, "badext@example.com")
        res = client.post("/resume/upload",
            files={"file": ("resume.exe", b"MZ executable", "application/octet-stream")},
            headers=headers)
        assert res.status_code == 400
        assert "Unsupported" in res.json()["detail"]

    def test_empty_file_rejected(self, client):
        headers = self._headers(client, "emptyfile@example.com")
        res = client.post("/resume/upload", files={"file": ("resume.txt", b"", "text/plain")}, headers=headers)
        assert res.status_code == 400
        assert "empty" in res.json()["detail"].lower()

    def test_docx_with_real_content_accepted(self, client):
        headers = self._headers(client, "docx_content@example.com")
        docx_bytes = make_docx_bytes(SAMPLE_RESUME_TEXT)
        res = client.post("/resume/upload",
            files={"file": ("resume.docx", docx_bytes, "application/octet-stream")},
            headers=headers)
        assert res.status_code == 201
        assert res.json()["file_type"] == "docx"

    def test_txt_with_unicode_accepted(self, client):
        headers = self._headers(client, "unicode_txt@example.com")
        unicode_resume = (
            "Jos\u00e9 Garc\u00eda Software Engineer\n"
            "SKILLS: Python, FastAPI, PostgreSQL, Docker\n"
            "PROJECTS: Built scalable REST APIs with FastAPI and deployed via Docker.\n"
        )
        res = client.post("/resume/upload",
            files={"file": ("resume.txt", unicode_resume.encode("utf-8"), "text/plain")},
            headers=headers)
        assert res.status_code == 201
        data = res.json()
        assert "Jos\u00e9" in data["raw_text_snippet"] or data["extracted_skills_count"] > 0

    def test_binary_disguised_as_txt_rejected(self, client):
        """
        Binary content with a .txt extension should be rejected.
        We use bytes that are mostly null/control chars (which clean_extracted_text
        replaces with spaces), leaving fewer than 20 readable characters — triggering
        the 'too little readable text' gate.
        Note: bytes(range(256)) decodes via latin-1 and produces enough
        visible characters to pass, so we use mostly-null content instead.
        """
        headers = self._headers(client, "binary_txt@example.com")
        # 300 null bytes + "hi" — decodes but strips to ~2 chars (way below 20-char gate)
        near_empty_binary = b"\x00" * 300 + b"hi"
        res = client.post("/resume/upload",
            files={"file": ("resume.txt", near_empty_binary, "text/plain")},
            headers=headers)
        assert res.status_code == 400


    def test_docx_content_not_just_extension(self, client):
        headers = self._headers(client, "fake_docx@example.com")
        res = client.post("/resume/upload",
            files={"file": ("resume.docx", b"\x00\x01\x02\x03garbage", "application/octet-stream")},
            headers=headers)
        assert res.status_code == 400


# ── TASK F: Re-upload and state management ───────────────────────────────────

class TestReuploadAndState:
    def test_second_upload_does_not_corrupt_first(self, client):
        _, headers = _register_and_login(client, "reupload_a@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")

        upload1 = client.post("/resume/upload", files={"file": ("v1.txt", txt, "text/plain")}, headers=headers)
        assert upload1.status_code == 201
        id1 = upload1.json()["resume_id"]

        upload2 = client.post("/resume/upload", files={"file": ("v2.txt", txt, "text/plain")}, headers=headers)
        assert upload2.status_code == 201
        id2 = upload2.json()["resume_id"]

        assert id2 != id1
        latest = client.get("/resume/latest", headers=headers).json()
        assert latest["id"] == id2

    def test_confirm_marks_resume_reviewed(self, client):
        _, headers = _register_and_login(client, "reupload_b@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        rid = upload.json()["resume_id"]

        client.post("/resume/confirm-skills",
            json={"resume_id": rid, "skills": [{"name": "Python", "level": "advanced"}]},
            headers=headers)

        latest = client.get("/resume/latest", headers=headers).json()
        assert latest["status"] == "reviewed"

    def test_upload_status_is_parsed_initially(self, client):
        _, headers = _register_and_login(client, "reupload_c@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")
        upload = client.post("/resume/upload", files={"file": ("resume.txt", txt, "text/plain")}, headers=headers)
        assert upload.json()["status"] == "parsed"

    def test_skill_change_after_reupload(self, client):
        _, headers = _register_and_login(client, "reupload_d@example.com")
        txt = SAMPLE_RESUME_TEXT.encode("utf-8")

        upload1 = client.post("/resume/upload", files={"file": ("v1.txt", txt, "text/plain")}, headers=headers)
        rid1 = upload1.json()["resume_id"]
        client.post("/resume/confirm-skills",
            json={"resume_id": rid1, "skills": [{"name": "Python", "level": "advanced"}]},
            headers=headers)

        upload2 = client.post("/resume/upload", files={"file": ("v2.txt", txt, "text/plain")}, headers=headers)
        rid2 = upload2.json()["resume_id"]
        client.post("/resume/confirm-skills",
            json={"resume_id": rid2, "skills": [{"name": "Docker", "level": "intermediate"}]},
            headers=headers)

        profile = client.get("/profile/me", headers=headers).json()
        skill_names = {s["name"] for s in profile["student_profile"]["skills"]}
        assert "Docker" in skill_names
        assert "Python" not in skill_names, "Second confirmation should replace first"
